import json
import tempfile
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.core.management import call_command
from django.test import TestCase
from django.db import IntegrityError
from django.contrib.auth.models import Permission
from django.urls import reverse
from docx import Document

from users.models import Organization, Role, User

from .models import Activity, ActivityIndicator, ActivityYearData, Commitment, CommitmentNarrativeReport, IcpdActualSubmission, IcpdExpenditureSubmission, IcpdSubmissionAudit, IndicatorYearData, Objective
from .models import CommitmentAccessPolicy
from .forms import IndicatorYearActualsForm, IndicatorYearDataForm


class IcpdPlanningModelTests(TestCase):
	def setUp(self):
		commitment = Commitment.objects.create(title='Improve adolescent health')
		objective = Objective.objects.create(
			commitment=commitment,
			title='Expand access to youth services',
		)
		self.activity = Activity.objects.create(
			objective=objective,
			title='Establish constituency innovation hubs',
			budget_amount=4000000000,
			responsibility='MoH',
		)
		self.activity_indicator = ActivityIndicator.objects.create(
			activity=self.activity,
			name='Innovation hubs established',
			baseline_value=0,
			baseline_year='2020/21',
		)

	def test_activity_has_one_budget_and_multiple_indicators(self):
		second_indicator = ActivityIndicator.objects.create(
			activity=self.activity,
			name='Young people reached',
		)

		self.assertEqual(self.activity.budget_amount, 4000000000)
		self.assertEqual(self.activity.activity_indicators.count(), 2)
		self.assertEqual(self.activity_indicator.code, f'ICPD-{self.activity_indicator.pk:04d}')
		self.assertEqual(second_indicator.code, f'ICPD-{second_indicator.pk:04d}')

	def test_only_one_yearly_record_per_indicator_and_financial_year(self):
		IndicatorYearData.objects.create(
			activity_indicator=self.activity_indicator,
			financial_year='2021/22',
			target_value=12,
			achievement_value=4,
			status='on_track',
		)

		with self.assertRaises(IntegrityError):
			IndicatorYearData.objects.create(
				activity_indicator=self.activity_indicator,
				financial_year='2021/22',
			)

	def test_cumulative_values_and_annual_activity_expenditure_are_tracked(self):
		IndicatorYearData.objects.create(
			activity_indicator=self.activity_indicator,
			financial_year='2021/22',
			target_value=12,
			achievement_value=4,
		)
		IndicatorYearData.objects.create(
			activity_indicator=self.activity_indicator,
			financial_year='2022/23',
			target_value=8,
			achievement_value=7,
		)
		ActivityYearData.objects.create(
			activity=self.activity,
			financial_year='2021/22',
			expenditure_amount=25,
		)

		self.assertEqual(self.activity_indicator.cumulative_target_value, 20)
		self.assertEqual(self.activity_indicator.cumulative_achievement_value, 11)

		with self.assertRaises(IntegrityError):
			ActivityYearData.objects.create(
				activity=self.activity,
				financial_year='2021/22',
			)

	def test_static_cumulative_target_uses_the_latest_annual_target(self):
		self.activity.cumulative_target_is_static = True
		self.activity.save(update_fields=['cumulative_target_is_static'])
		IndicatorYearData.objects.create(activity_indicator=self.activity_indicator, financial_year='2021/22', target_value='0.8')
		IndicatorYearData.objects.create(activity_indicator=self.activity_indicator, financial_year='2022/23', target_value='0.8')

		self.assertEqual(self.activity_indicator.cumulative_target_value, Decimal('0.8'))

	def test_contributor_form_exposes_only_annual_actuals_and_status(self):
		self.assertEqual(
			list(IndicatorYearActualsForm().fields),
			['achievement_value', 'status', 'remarks'],
		)

	def test_actual_submission_status_is_derived_from_target_attainment(self):
		status_for = IcpdActualSubmission.status_for_achievement

		self.assertEqual(status_for(None, 1), 'not_reported')
		self.assertEqual(status_for(0, 1), 'not_reported')
		self.assertEqual(status_for(10, 0), 'not_started')
		self.assertEqual(status_for(10, 7.4), 'at_risk')
		self.assertEqual(status_for(10, 7.5), 'on_track')
		self.assertEqual(status_for(10, 10), 'achieved')
		self.assertEqual(status_for(10, 12), 'achieved')

	def test_administrator_annual_target_form_excludes_actual_fields(self):
		self.assertEqual(
			list(IndicatorYearDataForm().fields),
			['activity_indicator', 'financial_year', 'target_value'],
		)

	def test_json_import_converts_direct_kes_budget_to_millions(self):
		payload = {
			'commitments': [{
				'title': 'Imported Commitment',
				'objectives': [{
					'title': 'Imported Objective',
					'activities': [{
						'title': 'Imported Activity',
						'budget_amount': 4000000000,
						'indicators': [{'name': 'Imported Indicator', 'annual_targets': []}],
					}],
				}],
			}],
		}
		with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as json_file:
			json.dump(payload, json_file)
			json_path = Path(json_file.name)
		try:
			call_command('import_icpd_plan', str(json_path))
		finally:
			json_path.unlink()

		activity = Activity.objects.get(title='Imported Activity')
		self.assertEqual(activity.budget_amount, 4000)
		self.assertEqual(activity.activity_indicators.get().code, 'ICPD-0002')

	def test_json_import_skips_existing_commitment(self):
		payload = {'commitments': [{'title': 'Existing Commitment'}]}
		Commitment.objects.create(title='Existing Commitment')
		with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as json_file:
			json.dump(payload, json_file)
			json_path = Path(json_file.name)
		try:
			call_command('import_icpd_plan', str(json_path))
		finally:
			json_path.unlink()

		self.assertEqual(Commitment.objects.filter(title='Existing Commitment').count(), 1)

	def test_json_import_accepts_source_key_action_schema(self):
		payload = {
			'commitments': [{
				'title': 'Imported Commitment',
				'objectives': [{
					'title': 'Imported Objective',
					'activities': [{
						'key_action': 'Imported key action',
						'indicator': 'Imported source indicator',
						'annual_values': {'2024/2025': 12, '2025/2026': '#####'},
						'remarks': None,
					}],
				}],
			}],
		}
		with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as json_file:
			json.dump(payload, json_file)
			json_path = Path(json_file.name)
		try:
			call_command('import_icpd_plan', str(json_path))
		finally:
			json_path.unlink()

		activity = Activity.objects.get(title='Imported key action')
		indicator = activity.activity_indicators.get(name='Imported source indicator')
		self.assertEqual(indicator.yearly_data.get().financial_year, '2024/25')
		self.assertEqual(indicator.yearly_data.get().target_value, 12)


class IcpdPermissionTests(TestCase):
	def test_role_with_add_commitment_permission_can_see_and_open_create_form(self):
		role = Role.objects.create(name='commitment_creator')
		role.permissions.add(
			Permission.objects.get(codename='view_commitment'),
			Permission.objects.get(codename='add_commitment'),
		)
		user = User.objects.create_user(
			username='commitment_creator', password='test-password', role=role,
		)
		self.client.force_login(user)

		dashboard_response = self.client.get(reverse('icpd:dashboard'))
		create_response = self.client.get(reverse('icpd:commitment_create'))

		self.assertContains(dashboard_response, reverse('icpd:commitment_create'))
		self.assertEqual(create_response.status_code, 200)

	def test_commitment_access_requires_its_own_view_permission(self):
		role = Role.objects.create(name='commitment_editor_without_access_policy')
		role.permissions.add(
			Permission.objects.get(codename='view_commitment'),
			Permission.objects.get(codename='change_commitment'),
		)
		user = User.objects.create_user(
			username='commitment_editor', password='test-password', role=role,
		)
		self.client.force_login(user)

		dashboard_response = self.client.get(reverse('icpd:dashboard'))
		access_response = self.client.get(reverse('icpd:commitment_access'))

		self.assertNotContains(dashboard_response, reverse('icpd:commitment_access'))
		self.assertRedirects(
			access_response,
			reverse('users:permission_denied'),
			fetch_redirect_response=False,
		)

	def test_duplicate_planning_entries_are_rejected_by_their_create_forms(self):
		user = User.objects.create_superuser(
			username='duplicate_planning_admin', email='duplicate-planning@example.com', password='test-password',
		)
		commitment = Commitment.objects.create(title='Duplicate Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Duplicate Objective')
		activity = Activity.objects.create(objective=objective, title='Duplicate Activity')
		ActivityIndicator.objects.create(activity=activity, name='Duplicate Indicator')
		self.client.force_login(user)

		responses = [
			self.client.post(reverse('icpd:commitment_create'), {
				'title': commitment.title, 'sort_order': 0,
			}),
			self.client.post(reverse('icpd:objective_create'), {
				'commitment': commitment.pk, 'title': objective.title, 'sort_order': 0,
			}),
			self.client.post(reverse('icpd:activity_create'), {
				'objective': objective.pk, 'title': activity.title, 'sort_order': 0,
			}),
			self.client.post(reverse('icpd:activity_indicator_create'), {
				'activity': activity.pk, 'name': 'Duplicate Indicator',
			}),
		]

		self.assertTrue(all(response.status_code == 200 for response in responses))
		self.assertEqual(Commitment.objects.filter(title=commitment.title).count(), 1)
		self.assertEqual(Objective.objects.filter(commitment=commitment, title=objective.title).count(), 1)
		self.assertEqual(Activity.objects.filter(objective=objective, title=activity.title).count(), 1)
		self.assertEqual(ActivityIndicator.objects.filter(activity=activity, name='Duplicate Indicator').count(), 1)

	def test_annual_target_create_excludes_already_recorded_financial_years(self):
		user = User.objects.create_superuser(
			username='annual_target_admin', email='annual-target@example.com', password='test-password',
		)
		commitment = Commitment.objects.create(title='Target Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Target Objective')
		activity = Activity.objects.create(objective=objective, title='Target Activity')
		indicator = ActivityIndicator.objects.create(activity=activity, name='Target Indicator')
		IndicatorYearData.objects.create(
			activity_indicator=indicator, financial_year='2025/26', target_value=10,
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:indicator_year_data_create'), {
			'activity_indicator': indicator.pk,
		})

		self.assertNotContains(response, '<option value="2025/26">', html=False)
		self.assertContains(response, '<option value="2026/27">', html=False)

	def test_activity_expenditure_create_redirects_to_submission_workflow(self):
		user = User.objects.create_superuser(
			username='expenditure_admin', email='expenditure@example.com', password='test-password',
		)
		commitment = Commitment.objects.create(title='Expenditure Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Expenditure Objective')
		activity = Activity.objects.create(objective=objective, title='Expenditure Activity')
		ActivityYearData.objects.create(
			activity=activity, financial_year='2025/26', expenditure_amount=10,
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:activity_year_data_create'), {
			'activity': activity.pk,
		})

		self.assertRedirects(response, reverse('icpd:report_entry'), fetch_redirect_response=False)

	def test_activity_indicator_create_url_resolves_for_an_icpd_administrator(self):
		user = User.objects.create_superuser(
			username='activity_indicator_admin',
			email='activity-indicator-admin@example.com',
			password='test-password',
		)
		commitment = Commitment.objects.create(title='Test Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Test Objective')
		activity = Activity.objects.create(objective=objective, title='Test Activity')
		self.client.force_login(user)

		response = self.client.get('/icpd/activity-indicators/add/', {'activity': activity.pk})

		self.assertEqual(response.status_code, 200)

	def test_dashboard_requires_view_commitment_permission(self):
		role = Role.objects.create(name='icpd_planner', display_name='ICPD Planner')
		user = User.objects.create_user(username='planner', password='test-password', role=role)

		self.client.force_login(user)
		denied_response = self.client.get('/icpd/')
		self.assertEqual(denied_response.status_code, 302)

		role.permissions.add(Permission.objects.get(codename='view_commitment'))
		allowed_response = self.client.get('/icpd/')
		self.assertEqual(allowed_response.status_code, 200)

	def test_objective_create_and_update_redirect_to_parent_commitment(self):
		user = User.objects.create_superuser(
			username='icpd_admin',
			email='icpd-admin@example.com',
			password='test-password',
		)
		commitment = Commitment.objects.create(title='Test Commitment')
		self.client.force_login(user)

		create_response = self.client.post(
			reverse('icpd:objective_create'),
			{'commitment': commitment.pk, 'title': 'Test Objective', 'sort_order': 1},
		)
		objective = Objective.objects.get(title='Test Objective')
		detail_url = reverse('icpd:commitment_detail', args=[commitment.pk])
		self.assertRedirects(create_response, detail_url)

		update_response = self.client.post(
			reverse('icpd:objective_update', args=[objective.pk]),
			{'commitment': commitment.pk, 'title': 'Updated Objective', 'sort_order': 1},
		)
		self.assertRedirects(update_response, detail_url)
		objective.refresh_from_db()
		self.assertEqual(objective.title, 'Updated Objective')

	def test_restricted_commitment_requires_an_assigned_user_or_role(self):
		commitment = Commitment.objects.create(title='Restricted Commitment')
		role = Role.objects.create(name='assigned_icpd_user')
		role.permissions.add(Permission.objects.get(codename='view_commitment'))
		user = User.objects.create_user(username='unassigned_icpd', password='test-password', role=role)
		policy = CommitmentAccessPolicy.objects.create(commitment=commitment, is_restricted=True)
		self.client.force_login(user)

		self.assertNotContains(self.client.get(reverse('icpd:dashboard')), 'Restricted Commitment')
		self.assertRedirects(self.client.get(reverse('icpd:commitment_detail', args=[commitment.pk])), reverse('icpd:dashboard'))

		policy.allowed_roles.add(role)
		self.assertContains(self.client.get(reverse('icpd:dashboard')), 'Restricted Commitment')
		self.assertEqual(self.client.get(reverse('icpd:commitment_detail', args=[commitment.pk])).status_code, 200)


class IcpdReportingFlowTests(TestCase):
	def test_report_entry_loads_without_a_reporting_selection(self):
		user = User.objects.create_superuser(
			username='empty_report_entry', email='empty-report@example.com', password='test-password',
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:report_entry'))

		self.assertRedirects(response, reverse('icpd:actual_report_list'))
		response = self.client.get(reverse('icpd:actual_report_list'))
		self.assertContains(response, 'Actual Reports')
		self.assertContains(response, 'Report Actuals')

	def test_assigned_user_can_save_a_commitment_narrative(self):
		commitment = Commitment.objects.create(title='Narrative Commitment')
		role = Role.objects.create(name='narrative_contributor')
		role.permissions.add(Permission.objects.get(codename='view_commitment'))
		user = User.objects.create_user(username='narrative_author', password='test-password', role=role)
		policy = CommitmentAccessPolicy.objects.create(commitment=commitment, is_restricted=True)
		policy.allowed_users.add(user)
		self.client.force_login(user)

		response = self.client.post(reverse('icpd:report_entry'), {
			'tab': 'narrative',
			'commitment': commitment.pk,
			'introduction': 'The commitment is on track.',
			'abbreviations': 'ICPD: International Conference on Population and Development.',
			'executive_summary': 'Implementation remains on track.',
			'other_actor_contributions': 'County partners supported delivery.',
			'facilitating_factors': 'Clear coordination enabled progress.',
			'challenges': 'Limited resources slowed progress.',
			'opportunities': 'New partnerships can enhance implementation.',
			'conclusion_and_recommendations': 'Continue county coordination.',
			'references': 'NCPD quarterly monitoring data, 2026.',
		})

		self.assertRedirects(response, f'{reverse("icpd:report_entry")}?tab=narrative&commitment={commitment.pk}&financial_year=2025/26')
		report = CommitmentNarrativeReport.objects.get(commitment=commitment, author=user)
		self.assertEqual(report.introduction, 'The commitment is on track.')
		self.assertEqual(report.workflow_status, 'submitted')

	def test_narrative_draft_can_be_updated_then_submitted_by_its_author(self):
		commitment = Commitment.objects.create(title='Draft Narrative Commitment')
		author = User.objects.create_superuser(username='draft_narrative_author', email='draft@example.com', password='test-password')
		self.client.force_login(author)
		response = self.client.post(reverse('icpd:report_entry'), {
			'tab': 'narrative', 'commitment': commitment.pk, 'financial_year': '2025/26',
			'introduction': 'First draft.', 'executive_summary': 'Draft summary.', 'submission_action': 'draft',
		})
		self.assertEqual(response.status_code, 302)
		narrative = CommitmentNarrativeReport.objects.get(commitment=commitment, financial_year='2025/26')
		self.assertEqual(narrative.workflow_status, 'draft')
		self.assertIsNone(narrative.submitted_at)
		response = self.client.post(reverse('icpd:report_entry'), {
			'tab': 'narrative', 'commitment': commitment.pk, 'financial_year': '2025/26',
			'introduction': 'Final narrative.', 'executive_summary': 'Final summary.', 'submission_action': 'submit',
		})
		self.assertEqual(response.status_code, 302)
		narrative.refresh_from_db()
		self.assertEqual(narrative.introduction, 'Final narrative.')
		self.assertEqual(narrative.workflow_status, 'submitted')
		self.assertIsNotNone(narrative.submitted_at)

	def test_submitted_narrative_is_locked_and_reviewer_can_approve_it(self):
		commitment = Commitment.objects.create(title='Locked Narrative Commitment')
		reporter = User.objects.create_superuser(username='narrative_reporter', email='narrative-reporter@example.com', password='test-password')
		narrative = CommitmentNarrativeReport.objects.create(commitment=commitment, author=reporter, financial_year='2025/26', introduction='Submitted narrative.')
		narrative.submit()
		narrative.save()
		self.client.force_login(reporter)
		response = self.client.post(reverse('icpd:report_entry'), {'tab': 'narrative', 'commitment': commitment.pk, 'financial_year': '2025/26', 'introduction': 'Changed narrative.'})
		self.assertEqual(response.status_code, 302)
		narrative.refresh_from_db()
		self.assertEqual(narrative.introduction, 'Submitted narrative.')
		reviewer = User.objects.create_superuser(username='narrative_reviewer', email='narrative-reviewer@example.com', password='test-password')
		self.client.force_login(reviewer)
		response = self.client.post(reverse('icpd:narrative_submission_review', args=[narrative.pk]), {'action': 'approve'})
		self.assertEqual(response.status_code, 302)
		narrative.refresh_from_db()
		self.assertEqual(narrative.workflow_status, 'approved')
		self.assertEqual(narrative.approved_by, reviewer)

	def test_narrative_export_is_hidden_until_a_narrative_is_saved(self):
		user = User.objects.create_superuser(username='empty_narrative_author', email='empty@example.com', password='test-password')
		commitment = Commitment.objects.create(title='Empty Narrative Commitment')
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:report_entry'), {
			'tab': 'narrative',
			'commitment': commitment.pk,
			'financial_year': '2025/26',
		})

		self.assertContains(response, 'Submit for Review')
		self.assertNotContains(response, 'Export Word')
		self.assertFalse(CommitmentNarrativeReport.objects.exists())

	def test_narrative_word_export_uses_the_selected_reporting_year(self):
		user = User.objects.create_superuser(username='narrative_exporter', email='export@example.com', password='test-password')
		commitment = Commitment.objects.create(title='Export Commitment')
		narrative = CommitmentNarrativeReport.objects.create(
			commitment=commitment,
			author=user,
			financial_year='2025/26',
			introduction='On track.',
			executive_summary='Strong progress.',
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:narrative_report_export'), {
			'commitment': commitment.pk,
			'financial_year': '2025/26',
			'narrative': narrative.pk,
		})

		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			response['Content-Type'],
			'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
		)
		document = Document(BytesIO(response.content))
		self.assertIn('ICPD COMMITMENT IMPLEMENTATION REPORT', [paragraph.text for paragraph in document.paragraphs])
		self.assertIn('Reported By', [cell.text for table in document.tables for row in table.rows for cell in row.cells])
	def test_report_entry_only_offers_the_selected_activity_indicator_years(self):
		user = User.objects.create_superuser(
			username='icpd_reporter', email='reporter@example.com', password='test-password',
		)
		commitment = Commitment.objects.create(title='Configured Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Configured Objective')
		activity = Activity.objects.create(objective=objective, title='Configured Activity')
		indicator = ActivityIndicator.objects.create(
			activity=activity, code='ICPD-CFG-01', name='Configured Activity Indicator',
		)
		IndicatorYearData.objects.create(
			activity_indicator=indicator, financial_year='2025/26', target_value=10,
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:report_entry'), {
			'commitment': commitment.pk,
			'objective': objective.pk,
			'activity': activity.pk,
			'activity_indicator': indicator.pk,
		})

		self.assertContains(response, 'Configured Activity Indicator')
		self.assertContains(response, '2025/26')
		self.assertNotContains(response, '2029/30')

	def test_report_entry_loads_actuals_form_for_selected_financial_year(self):
		organization = Organization.objects.create(name='Configured Organization')
		user = User.objects.create_superuser(
			username='year_selector', email='year-selector@example.com', password='test-password',
			organization=organization,
		)
		commitment = Commitment.objects.create(title='Configured Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Configured Objective')
		activity = Activity.objects.create(
			objective=objective, title='Configured Activity',
		)
		activity.responsible_organizations.add(organization)
		indicator = ActivityIndicator.objects.create(
			activity=activity, code='ICPD-CFG-SELECT', name='Configured Activity Indicator',
		)
		IndicatorYearData.objects.create(
			activity_indicator=indicator, financial_year='2025/26', target_value=10,
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:report_entry'), {
			'commitment': commitment.pk,
			'objective': objective.pk,
			'activity': activity.pk,
			'activity_indicator': indicator.pk,
			'financial_year': '2025/26',
		})

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Achievement Value')

	def test_reporter_creates_submissions_without_changing_the_configured_plan(self):
		organization = Organization.objects.create(name='Reporting Organization')
		user = User.objects.create_superuser(
			username='actuals_reporter', email='actuals@example.com', password='test-password',
			organization=organization,
		)
		commitment = Commitment.objects.create(title='Configured Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Configured Objective')
		activity = Activity.objects.create(objective=objective, title='Configured Activity')
		activity.responsible_organizations.add(organization)
		indicator = ActivityIndicator.objects.create(
			activity=activity, code='ICPD-CFG-02', name='Configured Activity Indicator',
		)
		year_data = IndicatorYearData.objects.create(
			activity_indicator=indicator, financial_year='2025/26', target_value=10,
		)
		self.client.force_login(user)

		response = self.client.post(reverse('icpd:report_entry'), {
			'commitment': commitment.pk,
			'objective': objective.pk,
			'activity': activity.pk,
			'activity_indicator': indicator.pk,
			'financial_year': '2025/26',
			'achievement_value': '7',
			'status': 'on_track',
			'expenditure-expenditure_amount': '25',
			'expenditure-remarks': 'Funds released.',
		})

		self.assertEqual(response.status_code, 302)
		year_data.refresh_from_db()
		self.assertIsNone(year_data.achievement_value)
		actual = IcpdActualSubmission.objects.get(indicator_year_data=year_data, submitted_by=user)
		expenditure = IcpdExpenditureSubmission.objects.get(activity=activity, submitted_by=user)
		self.assertEqual(actual.achievement_value, 7)
		self.assertEqual(expenditure.expenditure_amount, 25)
		self.assertEqual(actual.workflow_status, 'submitted')
		self.assertEqual(expenditure.workflow_status, 'submitted')

	def test_actual_report_detail_is_read_only(self):
		organization = Organization.objects.create(name='Actual Detail Organization')
		user = User.objects.create_superuser(username='actual_detail_user', email='actual-detail@example.com', password='test-password', organization=organization)
		commitment = Commitment.objects.create(title='Actual Detail Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Actual Detail Objective')
		activity = Activity.objects.create(objective=objective, title='Actual Detail Activity')
		activity.responsible_organizations.add(organization)
		indicator = ActivityIndicator.objects.create(activity=activity, name='Actual Detail Indicator')
		year_data = IndicatorYearData.objects.create(activity_indicator=indicator, financial_year='2025/26', target_value=10)
		report = IcpdExpenditureSubmission.objects.create(activity=activity, financial_year='2025/26', submitted_by=user, expenditure_amount=25, workflow_status='submitted')
		IcpdActualSubmission.objects.create(indicator_year_data=year_data, activity_report=report, submitted_by=user, achievement_value=7, status='on_track', workflow_status='submitted')
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:actual_report_detail', args=[report.pk]))

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Reported Actuals')
		self.assertNotContains(response, 'Submit for Review')

	def test_draft_reports_show_owner_edit_actions_and_approved_reports_show_exports(self):
		organization = Organization.objects.create(name='Detail Action Organization')
		user = User.objects.create_superuser(username='detail_action_user', email='detail-action@example.com', password='test-password', organization=organization)
		commitment = Commitment.objects.create(title='Detail Action Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Detail Action Objective')
		activity = Activity.objects.create(objective=objective, title='Detail Action Activity')
		activity.responsible_organizations.add(organization)
		indicator = ActivityIndicator.objects.create(activity=activity, name='Detail Action Indicator')
		year_data = IndicatorYearData.objects.create(activity_indicator=indicator, financial_year='2025/26', target_value=10)
		draft_report = IcpdExpenditureSubmission.objects.create(activity=activity, financial_year='2025/26', submitted_by=user, workflow_status='draft')
		IcpdActualSubmission.objects.create(indicator_year_data=year_data, activity_report=draft_report, submitted_by=user, workflow_status='draft')
		narrative = CommitmentNarrativeReport.objects.create(commitment=commitment, author=user, financial_year='2025/26', workflow_status='draft')
		self.client.force_login(user)

		self.assertContains(self.client.get(reverse('icpd:actual_report_detail', args=[draft_report.pk])), 'Edit Draft')
		self.assertContains(self.client.get(reverse('icpd:narrative_report_detail', args=[narrative.pk])), 'Edit Draft')
		draft_report.workflow_status = 'approved'
		draft_report.save(update_fields=['workflow_status'])
		narrative.workflow_status = 'approved'
		narrative.save(update_fields=['workflow_status'])
		self.assertContains(self.client.get(reverse('icpd:actual_report_detail', args=[draft_report.pk])), 'Excel')
		self.assertContains(self.client.get(reverse('icpd:narrative_report_detail', args=[narrative.pk])), 'PDF')
		response = self.client.get(reverse('icpd:narrative_detail_export', args=[narrative.pk, 'pdf']))
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response['Content-Type'], 'application/pdf')
		self.assertTrue(response.content.startswith(b'%PDF'))

	def test_pending_activity_one_submission_does_not_lock_activity_two(self):
		organization = Organization.objects.create(name='Responsible Organization')
		user = User.objects.create_superuser(
			username='separate_activity_reporter', email='separate@example.com', password='test-password',
			organization=organization,
		)
		commitment = Commitment.objects.create(title='Separate Activities Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Separate Activities Objective')
		first_activity = Activity.objects.create(objective=objective, title='First Activity')
		second_activity = Activity.objects.create(objective=objective, title='Second Activity')
		first_activity.responsible_organizations.add(organization)
		second_activity.responsible_organizations.add(organization)
		first_indicator = ActivityIndicator.objects.create(activity=first_activity, name='First Indicator')
		second_indicator = ActivityIndicator.objects.create(activity=second_activity, name='Second Indicator')
		first_year = IndicatorYearData.objects.create(
			activity_indicator=first_indicator, financial_year='2025/26', target_value=10,
		)
		IndicatorYearData.objects.create(
			activity_indicator=second_indicator, financial_year='2025/26', target_value=10,
		)
		first_report = IcpdExpenditureSubmission.objects.create(
			activity=first_activity, financial_year='2025/26', submitted_by=user,
			organization='Responsible Organization', workflow_status='submitted',
		)
		IcpdActualSubmission.objects.create(
			indicator_year_data=first_year, activity_report=first_report, submitted_by=user, workflow_status='submitted',
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:report_entry'), {
			'commitment': commitment.pk,
			'objective': objective.pk,
			'activity': second_activity.pk,
			'activity_indicator': second_indicator.pk,
			'financial_year': '2025/26',
		})

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Activity Report')

	def test_existing_submission_status_is_shown_for_the_selected_report(self):
		organization = Organization.objects.create(name='Responsible Organization')
		user = User.objects.create_superuser(
			username='submission_status_reporter', email='status@example.com', password='test-password',
			organization=organization,
		)
		commitment = Commitment.objects.create(title='Status Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Status Objective')
		activity = Activity.objects.create(objective=objective, title='Status Activity')
		activity.responsible_organizations.add(organization)
		indicator = ActivityIndicator.objects.create(activity=activity, name='Status Indicator')
		year_data = IndicatorYearData.objects.create(
			activity_indicator=indicator, financial_year='2025/26', target_value=10,
		)
		report = IcpdExpenditureSubmission.objects.create(
			activity=activity, financial_year='2025/26', submitted_by=user,
			organization='Responsible Organization', workflow_status='submitted',
		)
		IcpdActualSubmission.objects.create(
			indicator_year_data=year_data, activity_report=report, submitted_by=user, workflow_status='submitted',
		)
		self.client.force_login(user)

		response = self.client.get(reverse('icpd:report_entry'), {
			'commitment': commitment.pk, 'objective': objective.pk, 'activity': activity.pk,
			'activity_indicator': indicator.pk, 'financial_year': '2025/26',
		})

		self.assertContains(response, 'This activity already has a report for 2025/26.')
		self.assertContains(response, 'The activity is locked until a reviewer approves or returns the report.')
		self.assertContains(response, reverse('icpd:actual_report_detail', args=[report.pk]))

	def test_reviewer_approval_aggregates_submissions_and_blocks_contributor_overwrite(self):
		organization = Organization.objects.create(name='Reporting Organization')
		commitment = Commitment.objects.create(title='Aggregate Commitment')
		objective = Objective.objects.create(commitment=commitment, title='Aggregate Objective')
		activity = Activity.objects.create(objective=objective, title='Aggregate Activity')
		activity.responsible_organizations.add(organization)
		indicator = ActivityIndicator.objects.create(activity=activity, name='Aggregate Indicator')
		year_data = IndicatorYearData.objects.create(
			activity_indicator=indicator, financial_year='2025/26', target_value=10,
		)
		reporter_role = Role.objects.create(name='icpd_submission_reporter')
		reporter_role.permissions.add(
			Permission.objects.get(codename='view_commitment'),
			Permission.objects.get(codename='add_icpdactualsubmission'),
			Permission.objects.get(codename='add_icpdexpendituresubmission'),
		)
		reporter = User.objects.create_user(
			username='aggregate_reporter', password='test-password', role=reporter_role,
			organization=organization,
		)
		reviewer_role = Role.objects.create(name='icpd_submission_reviewer')
		reviewer_role.permissions.add(
			Permission.objects.get(codename='view_commitment'),
			Permission.objects.get(codename='can_approve_data'),
		)
		reviewer = User.objects.create_user(username='aggregate_reviewer', password='test-password', role=reviewer_role)
		self.client.force_login(reporter)

		response = self.client.post(reverse('icpd:report_entry'), {
			'commitment': commitment.pk, 'objective': objective.pk, 'activity': activity.pk,
			'activity_indicator': indicator.pk, 'financial_year': '2025/26',
			'achievement_value': '7', 'status': 'on_track',
			'expenditure-expenditure_amount': '25', 'expenditure-remarks': 'Released.',
		})
		self.assertEqual(response.status_code, 302)
		actual = IcpdActualSubmission.objects.get(indicator_year_data=year_data, submitted_by=reporter)
		expenditure = IcpdExpenditureSubmission.objects.get(activity=activity, submitted_by=reporter)
		self.client.force_login(reviewer)

		self.client.post(reverse('icpd:actual_submission_review', args=[actual.pk]), {'action': 'approve'})
		self.client.post(reverse('icpd:expenditure_submission_review', args=[expenditure.pk]), {'action': 'approve'})
		actual.refresh_from_db()
		expenditure.refresh_from_db()
		self.assertEqual(actual.workflow_status, 'approved')
		self.assertEqual(expenditure.workflow_status, 'approved')
		self.assertEqual(year_data.approved_achievement_value, 7)
		self.assertEqual(activity.approved_expenditure_for_year('2025/26'), 25)
		self.assertIsNone(year_data.achievement_value)
		self.assertEqual(IcpdSubmissionAudit.objects.filter(action='approved').count(), 2)

		response = self.client.post(
			reverse('icpd:actual_submission_review', args=[actual.pk]),
			{'action': 'reopen', 'reason': 'Please correct the source value.'},
		)
		self.assertEqual(response.status_code, 302)
		actual.refresh_from_db()
		self.assertEqual(actual.workflow_status, 'returned')
		self.assertEqual(actual.return_reason, 'Please correct the source value.')
