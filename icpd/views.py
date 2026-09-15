from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.conf import settings
from django.core.paginator import Paginator
from django.db.models import Count, Prefetch, Q, Sum
from django.utils import timezone
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from users.decorators import permission_required

from .models import Activity, ActivityIndicator, ActivityYearData, Commitment, CommitmentAccessPolicy, CommitmentNarrativeReport, FINANCIAL_YEAR_CHOICES, IcpdActualSubmission, IcpdExpenditureSubmission, IcpdSubmissionAudit, IndicatorYearData, Objective
from users.models import Role, User
from .forms import (
	ActivityForm,
	ActivityIndicatorForm,
	ActivityYearActualsForm,
	ActivityYearDataForm,
	CommitmentForm,
	CommitmentNarrativeReportForm,
	IndicatorYearActualsForm,
	IndicatorYearDataForm,
	IcpdActualSubmissionForm,
	IcpdExpenditureSubmissionForm,
	ObjectiveForm,
)


def _commitment_detail_url(instance):
	if isinstance(instance, Commitment):
		return 'icpd:commitment_detail', instance.pk
	if isinstance(instance, Objective):
		return 'icpd:commitment_detail', instance.commitment_id
	if isinstance(instance, Activity):
		return 'icpd:commitment_detail', instance.objective.commitment_id
	if isinstance(instance, ActivityIndicator):
		return 'icpd:commitment_detail', instance.activity.objective.commitment_id
	if isinstance(instance, ActivityYearData):
		return 'icpd:commitment_detail', instance.activity.objective.commitment_id
	return 'icpd:commitment_detail', instance.activity_indicator.activity.objective.commitment_id


def _user_can_access_commitment(user, commitment):
	"""Return whether a user is assigned to a commitment's optional access policy."""
	policy = getattr(commitment, 'access_policy', None)
	if not policy or not policy.is_restricted:
		return True
	return (
		policy.allowed_users.filter(pk=user.pk).exists()
		or (user.role_id and policy.allowed_roles.filter(pk=user.role_id).exists())
	)


def _visible_commitments(user):
	commitments = Commitment.objects.filter(is_active=True)
	return commitments.filter(
		Q(access_policy__isnull=True)
		| Q(access_policy__is_restricted=False)
		| Q(access_policy__allowed_users=user)
		| Q(access_policy__allowed_roles=user.role_id if user.role_id else None),
	).distinct()


def _rolling_financial_years(reporting_year):
	"""Return the reporting FY and six immediately preceding financial years."""
	start_year = int(reporting_year.split('/')[0])
	return [f'{year}/{str(year + 1)[-2:]}' for year in range(start_year - 6, start_year + 1)]


def _commitment_trend_rows(commitment, reporting_year):
	financial_years = _rolling_financial_years(reporting_year)[-5:]
	year_data = IndicatorYearData.objects.filter(
		activity_indicator__activity__objective__commitment=commitment,
		financial_year__in=financial_years,
	).select_related('activity_indicator').order_by('activity_indicator__code', 'financial_year')
	by_indicator = {}
	for item in year_data:
		by_indicator.setdefault(item.activity_indicator.code, {
			'indicator': item.activity_indicator.name,
			'values': {},
		})['values'][item.financial_year] = item
	return financial_years, by_indicator


def _commitment_trend_display_rows(commitment, reporting_year):
	financial_years, indicators = _commitment_trend_rows(commitment, reporting_year)
	return financial_years, [
		{
			'code': code,
			'indicator': data['indicator'],
			'targets': [data['values'].get(year).target_value if data['values'].get(year) else '-' for year in financial_years],
			'achievements': [data['values'].get(year).achievement_value if data['values'].get(year) else '-' for year in financial_years],
		}
		for code, data in indicators.items()
	]


def _narrative_has_content(narrative):
	return narrative and any((
		narrative.introduction,
		narrative.executive_summary,
		narrative.abbreviations,
		narrative.other_actor_contributions,
		narrative.facilitating_factors,
		narrative.challenges,
		narrative.opportunities,
		narrative.conclusion_and_recommendations,
		narrative.references,
	))


def _submission_values(submission, fields):
	return {field: str(getattr(submission, field) or '') for field in fields}


def _record_submission_audit(submission, submission_type, action, actor, old_value=None):
	fields = (
		('achievement_value', 'status', 'remarks', 'workflow_status')
		if submission_type == 'actual'
		else ('expenditure_amount', 'remarks', 'workflow_status')
	)
	IcpdSubmissionAudit.objects.create(
		submission_type=submission_type,
		submission_id=submission.pk,
		action=action,
		actor=actor,
		old_value=old_value,
		new_value=_submission_values(submission, fields),
	)


def _can_manage_actual_submission(user, submission):
	permission = 'change_icpdactualsubmission' if submission else 'add_icpdactualsubmission'
	return user.has_permission(permission) or user.has_permission('change_indicatoryeardata')


def _can_manage_expenditure_submission(user, submission):
	permission = 'change_icpdexpendituresubmission' if submission else 'add_icpdexpendituresubmission'
	return (
		user.has_permission(permission)
		or user.has_permission('add_activityyeardata')
		or user.has_permission('change_activityyeardata')
	)


def _user_is_activity_responsible(user, activity):
	return bool(
		user.organization_id
		and activity.responsible_organizations.filter(pk=user.organization_id).exists()
	)


def _require_commitment_access(request, commitment):
	if _user_can_access_commitment(request.user, commitment):
		return True
	messages.error(request, 'You are not assigned to this ICPD commitment.')
	return False


def _instance_commitment(instance):
	if isinstance(instance, Commitment):
		return instance
	if isinstance(instance, Objective):
		return instance.commitment
	if isinstance(instance, Activity):
		return instance.objective.commitment
	if isinstance(instance, ActivityIndicator):
		return instance.activity.objective.commitment
	if isinstance(instance, ActivityYearData):
		return instance.activity.objective.commitment
	return instance.activity_indicator.activity.objective.commitment


def _form_view(request, form_class, title, instance=None, initial=None, form_kwargs=None):
	if instance and not _require_commitment_access(request, _instance_commitment(instance)):
		return redirect('icpd:dashboard')
	form = form_class(
		request.POST or None,
		instance=instance,
		initial=initial,
		**(form_kwargs or {}),
	)
	if request.method == 'POST' and form.is_valid():
		saved_instance = form.save()
		messages.success(request, f'{title} saved successfully.')
		view_name, commitment_id = _commitment_detail_url(saved_instance)
		return redirect(view_name, pk=commitment_id)
	return render(request, 'icpd/form.html', {
		'form': form,
		'title': title,
		'is_edit': instance is not None,
	})


def _delete_view(request, instance, title):
	if not _require_commitment_access(request, _instance_commitment(instance)):
		return redirect('icpd:dashboard')
	if request.method == 'POST':
		view_name, commitment_id = _commitment_detail_url(instance)
		instance.delete()
		messages.success(request, f'{title} deleted successfully.')
		return redirect(view_name, pk=commitment_id)
	return render(request, 'icpd/confirm_delete.html', {'object': instance, 'title': title})


@login_required
@permission_required('view_commitment')
def dashboard(request):
	"""Show the ICPD commitments and a high-level delivery summary."""
	commitments = _visible_commitments(request.user).annotate(
		objective_count=Count('objectives', distinct=True),
		activity_count=Count('objectives__activities', distinct=True),
		indicator_count=Count(
			'objectives__activities__activity_indicators',
			distinct=True,
		),
	)
	context = {
		'commitments': commitments,
		'commitment_count': commitments.count(),
		'activity_count': Activity.objects.filter(objective__commitment__in=commitments).count(),
		'indicator_count': ActivityIndicator.objects.filter(activity__objective__commitment__in=commitments).count(),
		'on_track_count': IndicatorYearData.objects.filter(activity_indicator__activity__objective__commitment__in=commitments, status='on_track').count(),
		'at_risk_count': IndicatorYearData.objects.filter(activity_indicator__activity__objective__commitment__in=commitments, status='at_risk').count(),
		'can_add_commitment': request.user.has_permission('add_commitment'),
		'can_view_activity_indicators': request.user.has_permission('view_activityindicator'),
		'can_manage_access': request.user.has_permission('view_commitmentaccesspolicy'),
	}
	return render(request, 'icpd/dashboard.html', context)


@login_required
@permission_required('view_commitment')
def commitment_detail(request, pk):
	"""Display all objectives, activities, measures, and annual data for a commitment."""
	yearly_data = Prefetch(
		'yearly_data',
		queryset=IndicatorYearData.objects.order_by('financial_year'),
	)
	activity_indicators = Prefetch(
		'activity_indicators',
		queryset=ActivityIndicator.objects.prefetch_related(yearly_data),
	)
	activity_expenditure = Prefetch('yearly_expenditure', queryset=ActivityYearData.objects.order_by('financial_year'))
	commitment = get_object_or_404(
		Commitment.objects.prefetch_related(
		Prefetch('objectives__activities', queryset=Activity.objects.prefetch_related(activity_indicators, activity_expenditure)),
		),
		pk=pk,
	)
	if not _require_commitment_access(request, commitment):
		return redirect('icpd:dashboard')
	permissions = {
		'can_change_commitment': request.user.has_permission('change_commitment'),
		'can_delete_commitment': request.user.has_permission('delete_commitment'),
		'can_add_objective': request.user.has_permission('add_objective'),
		'can_change_objective': request.user.has_permission('change_objective'),
		'can_delete_objective': request.user.has_permission('delete_objective'),
		'can_add_activity': request.user.has_permission('add_activity'),
		'can_change_activity': request.user.has_permission('change_activity'),
		'can_delete_activity': request.user.has_permission('delete_activity'),
		'can_add_activity_indicator': request.user.has_permission('add_activityindicator'),
		'can_change_activity_indicator': request.user.has_permission('change_activityindicator'),
		'can_delete_activity_indicator': request.user.has_permission('delete_activityindicator'),
		'can_add_year_data': request.user.has_permission('add_indicatoryeardata'),
		'can_change_year_data': request.user.has_permission('change_indicatoryeardata'),
		'can_delete_year_data': request.user.has_permission('delete_indicatoryeardata'),
		'can_add_activity_year_data': request.user.has_permission('add_activityyeardata'),
		'can_change_activity_year_data': request.user.has_permission('change_activityyeardata'),
	}
	return render(request, 'icpd/commitment_detail.html', {
		'commitment': commitment,
		**permissions,
	})


@login_required
@permission_required('view_activityindicator')
def activity_indicator_list(request):
	"""Manage the ICPD-specific indicators assigned to planned activities."""
	indicators = ActivityIndicator.objects.select_related(
		'activity__objective__commitment',
	).prefetch_related('yearly_data')
	commitment_id = request.GET.get('commitment')
	objective_id = request.GET.get('objective')
	activity_id = request.GET.get('activity')
	if commitment_id:
		indicators = indicators.filter(activity__objective__commitment_id=commitment_id)
	if objective_id:
		indicators = indicators.filter(activity__objective_id=objective_id)
	if activity_id:
		indicators = indicators.filter(activity_id=activity_id)
	indicators = indicators.filter(activity__objective__commitment__in=_visible_commitments(request.user))
	paginator = Paginator(indicators, 25)
	page_obj = paginator.get_page(request.GET.get('page'))
	return render(request, 'icpd/activity_indicator_list.html', {
		'page_obj': page_obj,
		'activity_indicators': page_obj.object_list,
		'commitments': _visible_commitments(request.user),
		'objectives': Objective.objects.filter(commitment__in=_visible_commitments(request.user)).select_related('commitment'),
		'activities': Activity.objects.filter(objective__commitment__in=_visible_commitments(request.user)).select_related('objective__commitment'),
		'selected_commitment': commitment_id,
		'selected_objective': objective_id,
		'selected_activity': activity_id,
		'can_add_activity_indicator': request.user.has_permission('add_activityindicator'),
		'can_change_activity_indicator': request.user.has_permission('change_activityindicator'),
	})


@login_required
@permission_required('view_commitment')
def actual_report_list(request):
	commitments = _visible_commitments(request.user)
	activity_reports = IcpdExpenditureSubmission.objects.filter(
		activity__objective__commitment__in=commitments,
	).select_related('activity__objective__commitment', 'submitted_by').order_by('-updated_at')
	query = request.GET.get('q', '').strip()
	status = request.GET.get('status', '')
	financial_year = request.GET.get('financial_year', '')
	if query:
		activity_reports = activity_reports.filter(
			Q(activity__title__icontains=query)
			| Q(activity__objective__commitment__title__icontains=query)
			| Q(submitted_by__username__icontains=query)
		)
	if status:
		activity_reports = activity_reports.filter(workflow_status=status)
	if financial_year:
		activity_reports = activity_reports.filter(financial_year=financial_year)
	return render(request, 'icpd/actual_report_list.html', {
		'activity_reports': activity_reports,
		'financial_year_choices': FINANCIAL_YEAR_CHOICES,
		'selected_status': status,
		'selected_financial_year': financial_year,
		'search_query': query,
		'can_review_submissions': request.user.has_permission('can_approve_data'),
	})


@login_required
@permission_required('view_commitment')
def actual_report_detail(request, pk):
	report = get_object_or_404(
		IcpdExpenditureSubmission.objects.select_related(
			'activity__objective__commitment', 'submitted_by', 'approved_by',
		).prefetch_related(
			Prefetch('actual_submissions', queryset=IcpdActualSubmission.objects.select_related('indicator_year_data__activity_indicator')),
		),
		pk=pk,
	)
	if not _require_commitment_access(request, report.activity.objective.commitment):
		return redirect('icpd:dashboard')
	first_actual = report.actual_submissions.all()[0] if report.actual_submissions.all() else None
	return render(request, 'icpd/actual_report_detail.html', {
		'report': report,
		'can_edit': report.workflow_status == 'draft' and report.submitted_by_id == request.user.id,
		'edit_indicator_id': first_actual.indicator_year_data.activity_indicator_id if first_actual else None,
	})


@login_required
@permission_required('view_commitment')
def actual_report_export(request, pk, export_format):
	report = get_object_or_404(
		IcpdExpenditureSubmission.objects.select_related('activity__objective__commitment', 'submitted_by').prefetch_related(
			Prefetch('actual_submissions', queryset=IcpdActualSubmission.objects.select_related('indicator_year_data__activity_indicator')),
		), pk=pk,
	)
	if not _require_commitment_access(request, report.activity.objective.commitment):
		return redirect('icpd:actual_report_list')
	if report.workflow_status != 'approved' or export_format not in ('word', 'pdf', 'excel'):
		return HttpResponse('Only approved reports can be exported in this format.', status=403)
	rows = [[
		actual.indicator_year_data.activity_indicator.code,
		actual.indicator_year_data.activity_indicator.name,
		actual.indicator_year_data.target_value,
		actual.achievement_value,
		actual.get_status_display(),
		actual.remarks,
	] for actual in report.actual_submissions.all()]
	if export_format == 'excel':
		from openpyxl import Workbook
		from openpyxl.styles import Font, PatternFill
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = 'Actual Report'
		sheet.append(['ICPD Actual Report', report.activity.objective.commitment.title, report.activity.title, report.financial_year])
		sheet.append(['Reported by', report.submitted_by.get_full_name() or report.submitted_by.username])
		sheet.append([])
		headers = ['Code', 'Indicator', 'Target', 'Achievement', 'Status', 'Remarks']
		sheet.append(headers)
		for cell in sheet[4]:
			cell.font = Font(bold=True, color='FFFFFF')
			cell.fill = PatternFill('solid', fgColor='1A5632')
		for row in rows:
			sheet.append(row)
		for column in sheet.columns:
			sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or '')) for cell in column) + 2, 45)
		response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
		response['Content-Disposition'] = f'attachment; filename="icpd_actual_report_{report.pk}.xlsx"'
		workbook.save(response)
		return response
	if export_format == 'pdf':
		from reportlab.lib import colors
		from reportlab.lib.pagesizes import landscape, letter
		from reportlab.lib.styles import getSampleStyleSheet
		from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
		response = HttpResponse(content_type='application/pdf')
		response['Content-Disposition'] = f'attachment; filename="icpd_actual_report_{report.pk}.pdf"'
		styles = getSampleStyleSheet()
		content = [Paragraph('ICPD Actual Report', styles['Title']), Paragraph(f'{report.activity.objective.commitment.title} | {report.activity.title} | {report.financial_year}', styles['Normal']), Spacer(1, 12)]
		table = Table([['Code', 'Indicator', 'Target', 'Achievement', 'Status', 'Remarks']] + rows, repeatRows=1)
		table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A5632')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
		content.append(table)
		SimpleDocTemplate(response, pagesize=landscape(letter)).build(content)
		return response
	from docx import Document
	document = Document()
	document.add_heading('ICPD ACTUAL REPORT', 0)
	document.add_paragraph(f'Commitment: {report.activity.objective.commitment.title}')
	document.add_paragraph(f'Activity: {report.activity.title} | Financial Year: {report.financial_year}')
	document.add_paragraph(f'Reported by: {report.submitted_by.get_full_name() or report.submitted_by.username}')
	table = document.add_table(rows=1, cols=6)
	table.style = 'Light Grid Accent 1'
	for cell, value in zip(table.rows[0].cells, ['Code', 'Indicator', 'Target', 'Achievement', 'Status', 'Remarks']):
		cell.text = value
	for row in rows:
		for cell, value in zip(table.add_row().cells, row):
			cell.text = str(value or '-')
	response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
	response['Content-Disposition'] = f'attachment; filename="icpd_actual_report_{report.pk}.docx"'
	document.save(response)
	return response


@login_required
@permission_required('view_commitment')
def narrative_report_list(request):
	commitments = _visible_commitments(request.user)
	narratives = CommitmentNarrativeReport.objects.filter(
		commitment__in=commitments,
	).select_related('commitment', 'author').order_by('-updated_at')
	query = request.GET.get('q', '').strip()
	status = request.GET.get('status', '')
	financial_year = request.GET.get('financial_year', '')
	if query:
		narratives = narratives.filter(
			Q(commitment__title__icontains=query) | Q(author__username__icontains=query)
		)
	if status:
		narratives = narratives.filter(workflow_status=status)
	if financial_year:
		narratives = narratives.filter(financial_year=financial_year)
	return render(request, 'icpd/narrative_report_list.html', {
		'narratives': narratives,
		'financial_year_choices': FINANCIAL_YEAR_CHOICES,
		'selected_status': status,
		'selected_financial_year': financial_year,
		'search_query': query,
		'can_review_submissions': request.user.has_permission('can_approve_data'),
	})


@login_required
@permission_required('view_commitment')
def narrative_report_detail(request, pk):
	narrative = get_object_or_404(CommitmentNarrativeReport.objects.select_related('commitment', 'author', 'approved_by'), pk=pk)
	if not _require_commitment_access(request, narrative.commitment):
		return redirect('icpd:narrative_report_list')
	return render(request, 'icpd/narrative_report_detail.html', {
		'narrative': narrative,
		'can_edit': narrative.workflow_status == 'draft' and narrative.author_id == request.user.id,
		'narrative_sections': [
			('Introduction', narrative.introduction),
			('Executive Summary', narrative.executive_summary),
			('Abbreviations', narrative.abbreviations),
			('Contribution by Other Actors', narrative.other_actor_contributions),
			('Facilitating Factors', narrative.facilitating_factors),
			('Challenges', narrative.challenges),
			('Opportunities', narrative.opportunities),
			('Conclusion and Recommendations', narrative.conclusion_and_recommendations),
			('References', narrative.references),
		],
	})


@login_required
@permission_required('view_commitment')
def narrative_detail_export(request, pk, export_format):
	narrative = get_object_or_404(CommitmentNarrativeReport.objects.select_related('commitment', 'author'), pk=pk)
	if not _require_commitment_access(request, narrative.commitment):
		return redirect('icpd:narrative_report_list')
	if narrative.workflow_status != 'approved' or export_format not in ('word', 'pdf', 'excel'):
		return HttpResponse('Only approved reports can be exported in this format.', status=403)
	if export_format == 'word':
		return redirect(f'{request.path.rsplit(f"/narratives/{pk}/export/word/", 1)[0]}/narrative/export/?commitment={narrative.commitment_id}&financial_year={narrative.financial_year}&narrative={narrative.pk}')
	sections = [
		('Introduction', narrative.introduction), ('Executive Summary', narrative.executive_summary),
		('Abbreviations', narrative.abbreviations), ('Other Actor Contributions', narrative.other_actor_contributions),
		('Facilitating Factors', narrative.facilitating_factors), ('Challenges', narrative.challenges),
		('Opportunities', narrative.opportunities), ('Conclusion and Recommendations', narrative.conclusion_and_recommendations), ('References', narrative.references),
	]
	if export_format == 'excel':
		from openpyxl import Workbook
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = 'Narrative Report'
		sheet.append(['ICPD Commitment Narrative', narrative.commitment.title, narrative.financial_year])
		sheet.append(['Reported by', narrative.author.get_full_name() or narrative.author.username])
		for heading, content in sections:
			sheet.append([heading, content])
		response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
		response['Content-Disposition'] = f'attachment; filename="icpd_narrative_{narrative.pk}.xlsx"'
		workbook.save(response)
		return response
	from reportlab.lib import colors
	from reportlab.lib.enums import TA_CENTER
	from reportlab.lib.pagesizes import landscape, letter
	from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
	from reportlab.lib.units import inch
	from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
	from xml.sax.saxutils import escape
	response = HttpResponse(content_type='application/pdf')
	response['Content-Disposition'] = f'attachment; filename="icpd_narrative_{narrative.pk}.pdf"'
	styles = getSampleStyleSheet()
	title_style = ParagraphStyle('IcpdTitle', parent=styles['Title'], alignment=TA_CENTER, fontName='Helvetica-Bold', fontSize=16, leading=20)
	heading_style = ParagraphStyle('IcpdHeading', parent=styles['Heading2'], fontName='Helvetica-Bold', textColor=colors.HexColor('#1A5632'), spaceBefore=10, spaceAfter=5)
	body_style = ParagraphStyle('IcpdBody', parent=styles['BodyText'], fontSize=9, leading=13)
	metadata_style = ParagraphStyle('IcpdMetadata', parent=body_style, fontSize=8, leading=10)
	def pdf_text(value):
		return Paragraph(escape(str(value or 'Not provided.')).replace('\n', '<br/>'), body_style)
	def metadata_row(label, value):
		return [Paragraph(f'<b>{escape(label)}</b>', metadata_style), Paragraph(escape(str(value or '-')), metadata_style)]
	content = [Paragraph('ICPD COMMITMENT IMPLEMENTATION REPORT', title_style), Spacer(1, 0.15 * inch), Paragraph('Report Information', styles['Heading1'])]
	metadata = Table([
		metadata_row('Commitment', narrative.commitment.title),
		metadata_row('Reporting Financial Year', narrative.financial_year),
	], colWidths=[1.7 * inch, 8.4 * inch])
	metadata.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#B7C9BF')), ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#EAF2ED')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
	content.extend([metadata, Spacer(1, 0.18 * inch), Paragraph('Narrative Submission', styles['Heading1'])])
	submission_metadata = Table([
		metadata_row('Reported By', narrative.author.get_full_name() or narrative.author.username),
		metadata_row('Submitted On', timezone.localtime(narrative.submitted_at).strftime('%d %B %Y, %H:%M') if narrative.submitted_at else 'Not recorded'),
		metadata_row('Approval Status', narrative.get_workflow_status_display()),
		metadata_row('Approved On', timezone.localtime(narrative.approved_at).strftime('%d %B %Y, %H:%M') if narrative.approved_at else 'Pending approval'),
	], colWidths=[1.7 * inch, 8.4 * inch])
	submission_metadata.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#B7C9BF')), ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#EAF2ED')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
	content.extend([submission_metadata, Spacer(1, 0.15 * inch), Paragraph('a. Introduction (1-2 paragraphs)', heading_style), pdf_text(narrative.introduction), Paragraph('State the objectives, and planned actions for the selected financial year', heading_style)])
	for objective in Objective.objects.filter(commitment=narrative.commitment):
		content.append(Paragraph(f'• {escape(objective.title)}', body_style))
	for action in Activity.objects.filter(objective__commitment=narrative.commitment):
		content.append(Paragraph(f'    • {escape(action.title)}', body_style))
	content.extend([Paragraph('b. Executive summary', heading_style), pdf_text(narrative.executive_summary), Paragraph('c. Abbreviations', heading_style), pdf_text(narrative.abbreviations), Paragraph('d. Achievements (5 year period back from selected fy)', heading_style)])
	financial_years, trend_indicators = _commitment_trend_rows(narrative.commitment, narrative.financial_year)
	content.append(Paragraph(f'Trend period: {financial_years[0]} to {financial_years[-1]}', body_style))
	trend_rows = [[pdf_text('Indicator'), pdf_text('Measure')] + [pdf_text(year) for year in financial_years]]
	for data in trend_indicators.values():
		for measure in ('Target', 'Achievement'):
			row = [pdf_text(data['indicator']), pdf_text(measure)]
			for trend_year in financial_years:
				item = data['values'].get(trend_year)
				value = item.target_value if measure == 'Target' and item and item.target_value is not None else item.achievement_value if item and item.achievement_value is not None else '-'
				row.append(pdf_text(value))
			trend_rows.append(row)
	trend_table = Table(trend_rows, repeatRows=1, colWidths=[2.3 * inch, 0.75 * inch] + [0.68 * inch] * len(financial_years))
	trend_table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A5632')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#B7C9BF')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4), ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]))
	content.append(trend_table)
	for heading, value in sections:
		if heading not in ('Introduction', 'Executive Summary', 'Abbreviations'):
			content.extend([Paragraph(heading, heading_style), pdf_text(value), Spacer(1, 6)])
	SimpleDocTemplate(response, pagesize=landscape(letter), rightMargin=0.5 * inch, leftMargin=0.5 * inch, topMargin=0.5 * inch, bottomMargin=0.5 * inch).build(content)
	return response


@login_required
@permission_required('view_commitment')
def report_entry(request):
	"""Capture only annual ICPD actuals against an administrator-configured plan."""
	tab = request.GET.get('tab') or request.POST.get('tab') or 'actuals'
	commitment_id = request.GET.get('commitment') or request.POST.get('commitment')
	objective_id = request.GET.get('objective') or request.POST.get('objective')
	activity_id = request.GET.get('activity') or request.POST.get('activity')
	activity_indicator_id = request.GET.get('activity_indicator') or request.POST.get('activity_indicator')
	financial_year = request.GET.get('financial_year') or request.POST.get('financial_year')

	commitments = _visible_commitments(request.user)
	if tab == 'narrative':
		if not commitment_id and request.GET.get('mode') != 'entry':
			return redirect('icpd:narrative_report_list')
		financial_year = request.GET.get('financial_year') or request.POST.get('financial_year') or '2025/26'
		valid_financial_years = {value for value, _ in FINANCIAL_YEAR_CHOICES}
		if financial_year not in valid_financial_years:
			financial_year = '2025/26'
		selected_commitment = commitments.filter(pk=commitment_id).first() if commitment_id else None
		narrative_report = CommitmentNarrativeReport.objects.filter(
			commitment=selected_commitment,
			financial_year=financial_year,
		).first() if selected_commitment else None
		if request.method == 'POST' and selected_commitment:
			if narrative_report and (
				narrative_report.workflow_status != 'draft'
				or narrative_report.author_id != request.user.id
			):
				messages.error(request, 'This commitment narrative has already been submitted and is locked from editing.')
				return redirect(f'{request.path}?tab=narrative&commitment={selected_commitment.pk}&financial_year={financial_year}')
			form = CommitmentNarrativeReportForm(request.POST, instance=narrative_report)
			if form.is_valid():
				narrative = form.save(commit=False)
				narrative.commitment = selected_commitment
				narrative.author = request.user
				narrative.financial_year = financial_year
				action = request.POST.get('submission_action', 'submit')
				if action == 'submit':
					narrative.submit()
				else:
					narrative.workflow_status = 'draft'
				narrative.save()
				messages.success(request, 'Commitment narrative submitted for review.' if action == 'submit' else 'Commitment narrative saved as a draft.')
				return redirect(f'{request.path}?tab=narrative&commitment={selected_commitment.pk}&financial_year={financial_year}')
		else:
			form = CommitmentNarrativeReportForm(instance=narrative_report) if selected_commitment and (
				not narrative_report
				or (narrative_report.workflow_status == 'draft' and narrative_report.author_id == request.user.id)
			) else None
		trend_years, trend_rows = _commitment_trend_display_rows(selected_commitment, financial_year) if selected_commitment else ([], [])
		saved_narratives = CommitmentNarrativeReport.objects.filter(
			commitment=selected_commitment,
			financial_year=financial_year,
		).select_related('author') if selected_commitment else CommitmentNarrativeReport.objects.none()
		saved_narratives = [item for item in saved_narratives if _narrative_has_content(item)]
		return render(request, 'icpd/narrative_report.html', {
			'commitments': commitments,
			'selected_commitment': selected_commitment,
			'financial_year': financial_year,
			'financial_year_choices': FINANCIAL_YEAR_CHOICES,
			'form': form,
			'narrative_report': narrative_report,
			'can_submit_narrative': selected_commitment and (
				not narrative_report
				or (narrative_report.workflow_status == 'draft' and narrative_report.author_id == request.user.id)
			),
			'can_review_submissions': request.user.has_permission('can_approve_data'),
			'saved_narratives': saved_narratives,
			'planned_objectives': Objective.objects.filter(commitment=selected_commitment) if selected_commitment else Objective.objects.none(),
			'planned_actions': Activity.objects.filter(
				objective__commitment=selected_commitment,
			).select_related('objective') if selected_commitment else Activity.objects.none(),
			'trend_years': trend_years,
			'trend_rows': trend_rows,
		})
	if not commitment_id and request.GET.get('mode') != 'entry':
		return redirect('icpd:actual_report_list')
	objectives = Objective.objects.filter(commitment_id=commitment_id) if commitment_id else Objective.objects.none()
	activities = Activity.objects.filter(objective_id=objective_id) if objective_id else Activity.objects.none()
	activity_indicators = ActivityIndicator.objects.filter(activity_id=activity_id) if activity_id else ActivityIndicator.objects.none()
	if activity_id and not Activity.objects.filter(
		pk=activity_id,
		objective_id=objective_id,
		objective__commitment_id=commitment_id,
	).exists():
		messages.error(request, 'Select an activity that belongs to the selected objective and commitment.')
		return redirect('icpd:report_entry')
	if commitment_id and not commitments.filter(pk=commitment_id).exists():
		messages.error(request, 'You are not assigned to the selected ICPD commitment.')
		return redirect('icpd:report_entry')
	year_data = None
	actual_submission = None
	expenditure_submission = None
	financial_years = []
	if activity_indicator_id:
		financial_years = list(IndicatorYearData.objects.filter(
			activity_indicator_id=activity_indicator_id,
		).values_list('financial_year', 'financial_year'))
	if activity_indicator_id and financial_year:
		year_data = get_object_or_404(
			IndicatorYearData.objects.select_related('activity_indicator__activity__objective__commitment'),
			activity_indicator_id=activity_indicator_id,
			financial_year=financial_year,
		)
		if str(year_data.activity_indicator.activity_id) != str(activity_id):
			messages.error(request, 'The selected activity indicator does not belong to this activity.')
			return redirect('icpd:report_entry')
		expenditure_submission = IcpdExpenditureSubmission.objects.filter(
			activity_id=activity_id,
			financial_year=financial_year,
		).first()
		actual_submission = IcpdActualSubmission.objects.filter(
			indicator_year_data=year_data,
			activity_report=expenditure_submission,
		).first() if expenditure_submission else None

	if request.method == 'POST' and year_data:
		activity = year_data.activity_indicator.activity
		if not _user_is_activity_responsible(request.user, activity):
			messages.error(request, f'Only {activity.responsible_organization_names or "the assigned responsible organization"} can report actuals for this activity.')
			return redirect('icpd:report_entry')
		if expenditure_submission and expenditure_submission.workflow_status == 'draft' and expenditure_submission.submitted_by_id != request.user.id:
			messages.error(request, 'Only the reporter who created this draft can edit or submit it.')
			return redirect('icpd:actual_report_detail', pk=expenditure_submission.pk)
		if expenditure_submission and expenditure_submission.workflow_status == 'submitted':
			messages.error(request, 'This activity report is awaiting review and cannot be changed.')
			return redirect('icpd:report_entry')
		if expenditure_submission and expenditure_submission.workflow_status == 'approved':
			messages.error(request, 'This approved activity report is locked. Ask a reviewer to return it for correction.')
			return redirect('icpd:report_entry')
		if not (
			_can_manage_actual_submission(request.user, actual_submission)
			and _can_manage_expenditure_submission(request.user, expenditure_submission)
		):
			messages.error(request, 'You do not have permission to create or edit ICPD submissions.')
			return redirect('icpd:report_entry')
		actual_form = IcpdActualSubmissionForm(request.POST, instance=actual_submission)
		expenditure_form = IcpdExpenditureSubmissionForm(
			request.POST,
			instance=expenditure_submission or IcpdExpenditureSubmission(
				activity_id=activity_id,
				financial_year=financial_year,
				submitted_by=request.user,
				organization=request.user.organization.name if request.user.organization else '',
			),
			prefix='expenditure',
		)
		if actual_form.is_valid() and expenditure_form.is_valid():
			actual_old_value = _submission_values(actual_submission, ('achievement_value', 'status', 'remarks', 'workflow_status')) if actual_submission else None
			expenditure_old_value = _submission_values(expenditure_submission, ('expenditure_amount', 'remarks', 'workflow_status')) if expenditure_submission else None
			expenditure = expenditure_form.save(commit=False)
			expenditure.activity_id = activity_id
			expenditure.financial_year = financial_year
			if not expenditure.pk:
				expenditure.submitted_by = request.user
			expenditure.organization = request.user.organization.name if request.user.organization else ''
			action = request.POST.get('submission_action', 'submit')
			if action == 'submit':
				expenditure.submit()
			expenditure.save()
			actual = actual_form.save(commit=False)
			actual.indicator_year_data = year_data
			actual.activity_report = expenditure
			actual.status = IcpdActualSubmission.status_for_achievement(
				year_data.target_value,
				actual.achievement_value,
			)
			if not actual.pk:
				actual.submitted_by = request.user
			actual.organization = request.user.organization.name if request.user.organization else ''
			if action == 'submit':
				actual.submit()
			actual.save()
			audit_action = 'submitted' if action == 'submit' else ('updated' if actual_old_value else 'created')
			_record_submission_audit(actual, 'actual', audit_action, request.user, actual_old_value)
			_record_submission_audit(expenditure, 'expenditure', audit_action, request.user, expenditure_old_value)
			messages.success(request, 'ICPD submission sent for review.' if action == 'submit' else 'ICPD submission saved as a draft.')
			return redirect(f'{request.path}?commitment={commitment_id}&objective={objective_id}&activity={activity_id}&activity_indicator={activity_indicator_id}&financial_year={financial_year}')
	else:
		actual_form = IcpdActualSubmissionForm(instance=actual_submission) if year_data else None
		expenditure_form = IcpdExpenditureSubmissionForm(
			instance=expenditure_submission or IcpdExpenditureSubmission(
				activity_id=activity_id,
				financial_year=financial_year,
				submitted_by=request.user,
				organization=request.user.organization.name if request.user.organization else '',
			),
			prefix='expenditure',
		) if year_data else None

	return render(request, 'icpd/report_entry_workflow.html', {
		'commitments': commitments,
		'objectives': objectives,
		'activities': activities,
		'activity_indicators': activity_indicators,
		'financial_years': financial_years,
		'selected_commitment': commitment_id,
		'selected_objective': objective_id,
		'selected_activity': activity_id,
		'selected_activity_indicator': activity_indicator_id,
		'selected_financial_year': financial_year,
		'year_data': year_data,
		'actual_submission': actual_submission,
		'expenditure_submission': expenditure_submission,
		'activity_is_responsible': _user_is_activity_responsible(request.user, year_data.activity_indicator.activity) if year_data else False,
		'approved_achievement_value': year_data.approved_achievement_value if year_data else None,
		'approved_expenditure_amount': Activity.objects.get(pk=activity_id).approved_expenditure_for_year(financial_year) if year_data else None,
		'can_review_submissions': request.user.has_permission('can_approve_data'),
		'year_form': actual_form,
		'expenditure_form': expenditure_form,
		'can_edit_submission': (
			year_data
			and (not expenditure_submission or expenditure_submission.workflow_status in ('draft', 'returned'))
			and _user_is_activity_responsible(request.user, year_data.activity_indicator.activity)
			and _can_manage_actual_submission(request.user, actual_submission)
			and _can_manage_expenditure_submission(request.user, expenditure_submission)
		),
	})


def _review_submission(request, submission, submission_type):
	if not _require_commitment_access(request, _instance_commitment(
		submission.indicator_year_data if submission_type == 'actual' else submission.activity,
	)):
		return redirect('icpd:dashboard')
	if submission.submitted_by_id == request.user.id:
		messages.error(request, 'You cannot review your own ICPD submission.')
		return redirect('icpd:submission_review')
	action = request.POST.get('action')
	if submission.workflow_status != 'submitted' and not (
		submission.workflow_status == 'approved' and action == 'reopen'
	):
		messages.error(request, 'Only submitted ICPD records can be reviewed.')
		return redirect('icpd:submission_review')

	if action == 'approve':
		old_value = _submission_values(submission, ('achievement_value', 'status', 'remarks', 'workflow_status') if submission_type == 'actual' else ('expenditure_amount', 'remarks', 'workflow_status'))
		submission.workflow_status = 'approved'
		submission.approved_by = request.user
		submission.approved_at = timezone.now()
		submission.save()
		_record_submission_audit(submission, submission_type, 'approved', request.user, old_value)
		messages.success(request, 'ICPD submission approved and included in the official aggregate.')
	elif action in ('return', 'reopen'):
		reason = request.POST.get('reason', '').strip()
		if not reason:
			messages.error(request, 'Provide a reason when returning an ICPD submission for correction.')
			return redirect('icpd:submission_review')
		old_value = _submission_values(submission, ('achievement_value', 'status', 'remarks', 'workflow_status') if submission_type == 'actual' else ('expenditure_amount', 'remarks', 'workflow_status'))
		submission.workflow_status = 'returned'
		submission.returned_by = request.user
		submission.returned_at = timezone.now()
		submission.return_reason = reason
		submission.save()
		_record_submission_audit(submission, submission_type, 'returned', request.user, old_value)
		messages.warning(request, 'ICPD submission returned to its contributor for correction.')
	else:
		messages.error(request, 'Unknown review action.')
	return redirect('icpd:submission_review')


@login_required
@permission_required('can_approve_data')
def submission_review(request):
	"""Review ICPD contributor submissions before they form official totals."""
	visible = _visible_commitments(request.user)
	actual_submissions = IcpdActualSubmission.objects.filter(
		workflow_status='submitted',
		indicator_year_data__activity_indicator__activity__objective__commitment__in=visible,
	).select_related('indicator_year_data__activity_indicator__activity__objective__commitment', 'submitted_by')
	expenditure_submissions = IcpdExpenditureSubmission.objects.filter(
		workflow_status='submitted',
		activity__objective__commitment__in=visible,
	).select_related('activity__objective__commitment', 'submitted_by')
	approved_actual_submissions = IcpdActualSubmission.objects.filter(
		workflow_status='approved',
		indicator_year_data__activity_indicator__activity__objective__commitment__in=visible,
	).select_related('indicator_year_data__activity_indicator__activity__objective__commitment', 'submitted_by')
	approved_expenditure_submissions = IcpdExpenditureSubmission.objects.filter(
		workflow_status='approved',
		activity__objective__commitment__in=visible,
	).select_related('activity__objective__commitment', 'submitted_by')
	narrative_submissions = CommitmentNarrativeReport.objects.filter(
		workflow_status='submitted', commitment__in=visible,
	).select_related('commitment', 'author')
	approved_narrative_submissions = CommitmentNarrativeReport.objects.filter(
		workflow_status='approved', commitment__in=visible,
	).select_related('commitment', 'author', 'approved_by')
	return render(request, 'icpd/submission_review.html', {
		'actual_submissions': actual_submissions,
		'expenditure_submissions': expenditure_submissions,
		'approved_actual_submissions': approved_actual_submissions,
		'approved_expenditure_submissions': approved_expenditure_submissions,
		'narrative_submissions': narrative_submissions,
		'approved_narrative_submissions': approved_narrative_submissions,
	})


@login_required
@permission_required('can_approve_data')
def actual_submission_review(request, pk):
	if request.method != 'POST':
		return redirect('icpd:submission_review')
	return _review_submission(request, get_object_or_404(IcpdActualSubmission, pk=pk), 'actual')


@login_required
@permission_required('can_approve_data')
def expenditure_submission_review(request, pk):
	if request.method != 'POST':
		return redirect('icpd:submission_review')
	return _review_submission(request, get_object_or_404(IcpdExpenditureSubmission, pk=pk), 'expenditure')


@login_required
@permission_required('can_approve_data')
def narrative_submission_review(request, pk):
	if request.method != 'POST':
		return redirect('icpd:submission_review')
	narrative = get_object_or_404(CommitmentNarrativeReport, pk=pk)
	if not _require_commitment_access(request, narrative.commitment):
		return redirect('icpd:dashboard')
	if narrative.author_id == request.user.id:
		messages.error(request, 'You cannot review your own ICPD narrative submission.')
		return redirect('icpd:submission_review')
	if narrative.workflow_status != 'submitted':
		messages.error(request, 'Only submitted commitment narratives can be approved.')
		return redirect('icpd:submission_review')
	if request.POST.get('action') != 'approve':
		messages.error(request, 'Unknown review action.')
		return redirect('icpd:submission_review')
	narrative.workflow_status = 'approved'
	narrative.approved_by = request.user
	narrative.approved_at = timezone.now()
	narrative.save(update_fields=['workflow_status', 'approved_by', 'approved_at', 'updated_at'])
	messages.success(request, 'ICPD commitment narrative approved.')
	return redirect('icpd:submission_review')


@login_required
@permission_required('view_commitment')
def narrative_report_export(request):
	"""Export an ICPD commitment narrative and its rolling trend as Word."""
	from docx import Document
	from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
	from docx.shared import Inches, Pt

	commitment = get_object_or_404(Commitment, pk=request.GET.get('commitment'))
	financial_year = request.GET.get('financial_year', '2025/26')
	if financial_year not in {value for value, _ in FINANCIAL_YEAR_CHOICES}:
		return HttpResponse('Unknown reporting financial year.', status=400)
	if not _require_commitment_access(request, commitment):
		return redirect('icpd:report_entry')

	narratives = CommitmentNarrativeReport.objects.filter(
		commitment=commitment,
		financial_year=financial_year,
	).select_related('author', 'approved_by')
	selected_narrative_ids = request.GET.getlist('narrative')
	if not selected_narrative_ids:
		messages.error(request, 'Select at least one saved narrative to export.')
		return redirect(f'{request.path.rsplit("/export/", 1)[0]}/?tab=narrative&commitment={commitment.pk}&financial_year={financial_year}')
	narratives = narratives.filter(pk__in=selected_narrative_ids).exclude(workflow_status='draft')
	if not any(_narrative_has_content(narrative) for narrative in narratives):
		messages.error(request, 'Submit narrative content before exporting this report.')
		return redirect(f'{request.path.rsplit("/export/", 1)[0]}/?tab=narrative&commitment={commitment.pk}&financial_year={financial_year}')

	document = Document()
	logo_path = settings.BASE_DIR / 'static' / 'images' / 'logo.png'
	if logo_path.exists():
		logo = document.add_picture(str(logo_path), width=Inches(2.0))
		logo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
	styles = document.styles
	for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
		style = styles[style_name]
		style.font.name = 'Tahoma'
		style.font.size = Pt(12)
		style.paragraph_format.line_spacing = 1.5

	title = document.add_heading('ICPD COMMITMENT IMPLEMENTATION REPORT', level=0)
	title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
	document.add_heading('Report Information', level=1)
	metadata = document.add_table(rows=2, cols=2)
	metadata.style = 'Light Grid Accent 1'
	metadata.rows[0].cells[0].text = 'Commitment'
	metadata.rows[0].cells[1].text = commitment.title
	metadata.rows[1].cells[0].text = 'Reporting Financial Year'
	metadata.rows[1].cells[1].text = financial_year
	financial_years, trend_indicators = _commitment_trend_rows(commitment, financial_year)
	planned_objectives = Objective.objects.filter(commitment=commitment)
	planned_actions = Activity.objects.filter(
		objective__commitment=commitment,
	).select_related('objective')
	for narrative in narratives:
		document.add_heading('Narrative Submission', level=1)
		submission_metadata = document.add_table(rows=4, cols=2)
		submission_metadata.style = 'Light Grid Accent 1'
		submission_metadata.rows[0].cells[0].text = 'Reported By'
		submission_metadata.rows[0].cells[1].text = narrative.author.get_full_name() or narrative.author.username
		submission_metadata.rows[1].cells[0].text = 'Submitted On'
		submission_metadata.rows[1].cells[1].text = timezone.localtime(narrative.submitted_at).strftime('%d %B %Y, %H:%M') if narrative.submitted_at else 'Not recorded'
		submission_metadata.rows[2].cells[0].text = 'Approval Status'
		submission_metadata.rows[2].cells[1].text = narrative.get_workflow_status_display()
		submission_metadata.rows[3].cells[0].text = 'Approved On'
		submission_metadata.rows[3].cells[1].text = timezone.localtime(narrative.approved_at).strftime('%d %B %Y, %H:%M') if narrative.approved_at else 'Pending approval'
		document.add_heading('a. Introduction (1-2 paragraphs)', level=2)
		document.add_paragraph(narrative.introduction or 'Not provided.')
		document.add_heading('State the objectives, and planned actions for the selected financial year', level=2)
		for objective in planned_objectives:
			document.add_paragraph(objective.title, style='List Bullet')
		for action in planned_actions:
			document.add_paragraph(action.title, style='List Bullet 2')
		document.add_heading('b. Executive summary', level=2)
		document.add_paragraph(narrative.executive_summary or 'Not provided.')
		document.add_heading('c. Abbreviations', level=2)
		document.add_paragraph(narrative.abbreviations or 'Not provided.')
		document.add_heading('d. Achievements (5 year period back from selected fy)', level=2)
		document.add_paragraph(f'Trend period: {financial_years[0]} to {financial_years[-1]}')
		table = document.add_table(rows=1, cols=2 + len(financial_years))
		table.style = 'Light Grid Accent 1'
		for index, value in enumerate(['Indicator', 'Measure'] + financial_years):
			table.rows[0].cells[index].text = value
		for data in trend_indicators.values():
			for measure in ('Target', 'Achievement'):
				cells = table.add_row().cells
				cells[0].text = data['indicator']
				cells[1].text = measure
				for index, trend_year in enumerate(financial_years, start=2):
					item = data['values'].get(trend_year)
					cells[index].text = str(item.target_value if measure == 'Target' and item and item.target_value is not None else item.achievement_value if item and item.achievement_value is not None else '-')
		for heading, content in (
			('e. Contribution by other actors', narrative.other_actor_contributions),
			('f. Facilitating factors', narrative.facilitating_factors),
			('What challenges slowed progress?', narrative.challenges),
			('What opportunities to enhance implementation exist?', narrative.opportunities),
			('g. Conclusion and Recommendations', narrative.conclusion_and_recommendations),
			('h. References', narrative.references),
		):
			document.add_heading(heading, level=2)
			document.add_paragraph(content or 'Not provided.')

	response = HttpResponse(
		content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
	)
	response['Content-Disposition'] = f'attachment; filename="icpd_commitment_{commitment.pk}_{financial_year.replace("/", "_")}.docx"'
	document.save(response)
	return response


@login_required
@permission_required('add_commitment')
def commitment_create(request):
	return _form_view(request, CommitmentForm, 'Commitment')


@login_required
@permission_required('change_commitment')
def commitment_update(request, pk):
	return _form_view(request, CommitmentForm, 'Commitment', get_object_or_404(Commitment, pk=pk))


@login_required
@permission_required('view_commitmentaccesspolicy')
def commitment_access(request):
	"""Assign roles and users that may work on each ICPD commitment."""
	commitments = Commitment.objects.filter(is_active=True)
	if request.method == 'POST':
		roles = Role.objects.filter(is_active=True)
		users = User.objects.filter(is_active=True, is_verified=True)
		for commitment in commitments:
			policy = CommitmentAccessPolicy.objects.filter(commitment=commitment).first()
			permission = 'change_commitmentaccesspolicy' if policy else 'add_commitmentaccesspolicy'
			if not request.user.has_permission(permission):
				messages.error(request, 'You do not have permission to update ICPD commitment access settings.')
				return redirect('icpd:commitment_access')
			if policy is None:
				policy = CommitmentAccessPolicy.objects.create(commitment=commitment)
			policy.is_restricted = request.POST.get(f'restricted_{commitment.pk}') == 'on'
			policy.save()
			policy.allowed_roles.set(roles.filter(pk__in=request.POST.getlist(f'roles_{commitment.pk}')))
			policy.allowed_users.set(users.filter(pk__in=request.POST.getlist(f'users_{commitment.pk}')))
		messages.success(request, 'ICPD commitment access settings updated.')
		return redirect('icpd:commitment_access')
	policies = {
		policy.commitment_id: policy
		for policy in CommitmentAccessPolicy.objects.prefetch_related('allowed_roles', 'allowed_users')
	}
	settings = []
	for commitment in commitments:
		policy = policies.get(commitment.pk)
		settings.append({
			'commitment': commitment,
			'is_restricted': policy.is_restricted if policy else False,
			'allowed_role_ids': list(policy.allowed_roles.values_list('id', flat=True)) if policy else [],
			'allowed_user_ids': list(policy.allowed_users.values_list('id', flat=True)) if policy else [],
		})
	return render(request, 'icpd/access.html', {
		'commitment_access_settings': settings,
		'roles': Role.objects.filter(is_active=True).order_by('display_name', 'name'),
		'users': User.objects.filter(is_active=True, is_verified=True).select_related('role').order_by('username'),
	})


@login_required
@permission_required('delete_commitment')
def commitment_delete(request, pk):
	return _delete_view(request, get_object_or_404(Commitment, pk=pk), 'Commitment')


@login_required
@permission_required('add_objective')
def objective_create(request):
	commitment = get_object_or_404(Commitment, pk=request.GET.get('commitment') or request.POST.get('commitment'))
	if not _require_commitment_access(request, commitment):
		return redirect('icpd:dashboard')
	return _form_view(request, ObjectiveForm, 'Objective', initial={'commitment': commitment.pk})


@login_required
@permission_required('change_objective')
def objective_update(request, pk):
	return _form_view(request, ObjectiveForm, 'Objective', get_object_or_404(Objective, pk=pk))


@login_required
@permission_required('delete_objective')
def objective_delete(request, pk):
	return _delete_view(request, get_object_or_404(Objective, pk=pk), 'Objective')


@login_required
@permission_required('add_activity')
def activity_create(request):
	objective = get_object_or_404(Objective.objects.select_related('commitment'), pk=request.GET.get('objective') or request.POST.get('objective'))
	if not _require_commitment_access(request, objective.commitment):
		return redirect('icpd:dashboard')
	return _form_view(request, ActivityForm, 'Activity', initial={'objective': objective.pk})


@login_required
@permission_required('change_activity')
def activity_update(request, pk):
	return _form_view(request, ActivityForm, 'Activity', get_object_or_404(Activity, pk=pk))


@login_required
@permission_required('delete_activity')
def activity_delete(request, pk):
	return _delete_view(request, get_object_or_404(Activity, pk=pk), 'Activity')


@login_required
@permission_required('add_activityindicator')
def activity_indicator_create(request):
	activity_id = request.GET.get('activity') or request.POST.get('activity')
	if not activity_id:
		return _form_view(request, ActivityIndicatorForm, 'Activity Indicator')
	activity = get_object_or_404(Activity.objects.select_related('objective__commitment'), pk=activity_id)
	if not _require_commitment_access(request, activity.objective.commitment):
		return redirect('icpd:dashboard')
	return _form_view(request, ActivityIndicatorForm, 'Activity Indicator', initial={'activity': activity.pk})


@login_required
@permission_required('change_activityindicator')
def activity_indicator_update(request, pk):
	return _form_view(request, ActivityIndicatorForm, 'Activity Indicator', get_object_or_404(ActivityIndicator, pk=pk))


@login_required
@permission_required('delete_activityindicator')
def activity_indicator_delete(request, pk):
	return _delete_view(request, get_object_or_404(ActivityIndicator, pk=pk), 'Activity Indicator')


@login_required
@permission_required('add_indicatoryeardata')
def indicator_year_data_create(request):
	indicator = get_object_or_404(
		ActivityIndicator.objects.select_related('activity__objective__commitment'),
		pk=request.GET.get('activity_indicator') or request.POST.get('activity_indicator'),
	)
	if not _require_commitment_access(request, indicator.activity.objective.commitment):
		return redirect('icpd:dashboard')
	if IndicatorYearData.objects.filter(activity_indicator=indicator).count() >= len(FINANCIAL_YEAR_CHOICES):
		messages.info(request, 'All financial years already have annual targets. Use the edit action to update an existing target.')
		return redirect('icpd:commitment_detail', pk=indicator.activity.objective.commitment_id)
	return _form_view(
		request,
		IndicatorYearDataForm,
		'Annual Target',
		initial={'activity_indicator': indicator.pk},
		form_kwargs={'activity_indicator': indicator},
	)


@login_required
@permission_required('change_indicatoryeardata')
def indicator_year_data_update(request, pk):
	year_data = get_object_or_404(IndicatorYearData, pk=pk)
	return _form_view(request, IndicatorYearDataForm, 'Annual Target', year_data)


@login_required
@permission_required('delete_indicatoryeardata')
def indicator_year_data_delete(request, pk):
	return _delete_view(request, get_object_or_404(IndicatorYearData, pk=pk), 'Yearly Indicator Data')


@login_required
@permission_required('add_activityyeardata')
def activity_year_data_create(request):
	messages.info(request, 'Activity expenditure is reported through ICPD submissions and reviewed before it is official.')
	return redirect('icpd:report_entry')


@login_required
@permission_required('change_activityyeardata')
def activity_year_data_update(request, pk):
	messages.info(request, 'Activity expenditure is reported through ICPD submissions and reviewed before it is official.')
	return redirect('icpd:report_entry')

# Create your views here.
